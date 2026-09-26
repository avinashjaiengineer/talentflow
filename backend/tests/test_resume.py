import io
from datetime import datetime
from pathlib import Path

from app.resume import (
    Position,
    _heuristic_profile,
    extract_text,
    guess_name,
    parse_month,
    split_sections,
    years_from_history,
)

TEST_DATA = Path(__file__).resolve().parents[2] / "docs" / "test-data"
TODAY = datetime(2026, 9, 1)


def test_parse_month_understands_common_resume_dates():
    assert parse_month("2019-03") == (2019, 3)
    assert parse_month("03/2019") == (2019, 3)
    assert parse_month("Mar 2019") == (2019, 3)
    assert parse_month("September 2021") == (2021, 9)
    assert parse_month("Jan '22") == (2022, 1)
    assert parse_month("2015") == (2015, 1)
    assert parse_month("Present", today=TODAY) == (2026, 9)
    assert parse_month("till date", today=TODAY) == (2026, 9)
    assert parse_month("sometime") is None
    assert parse_month("13/2019") is None


def test_years_from_history_counts_overlapping_jobs_once():
    history = [
        Position(title="Staff Engineer", start="2019-01", end="present"),
        Position(title="Consultant", start="2020-01", end="2021-01"),  # overlaps the job above
        Position(title="Engineer", start="2015-01", end="2018-01"),
        Position(title="Undated"),
    ]
    assert years_from_history(history, today=datetime(2025, 1, 1)) == 9.0
    assert years_from_history([Position(title="x")]) is None


def test_sections_include_inline_headings():
    sections = split_sections("Jane Doe\nSkills: Python, Go\nEXPERIENCE\nAcme (2020 - 2022), Engineer\nEducation:\nBSc CS, 2019")
    assert sections["skills"] == ["Python, Go"]
    assert sections["experience"] == ["Acme (2020 - 2022), Engineer"]
    assert sections["education"] == ["BSc CS, 2019"]


def test_guess_name_skips_headings_contacts_and_titles():
    lines = ["RESUME", "john@example.com | +1 415 555 0100", "John Smith | Senior Data Engineer", "Summary"]
    assert guess_name(lines) == "John Smith"
    assert guess_name(["Curriculum Vitae", "Ananya Iyer - Product Designer"]) == "Ananya Iyer"


def test_offline_parser_extracts_history_skills_and_links():
    text = (TEST_DATA / "01-strong-backend-fit.txt").read_text(encoding="utf-8") + "\nlinkedin.com/in/arjunmehta"
    p = _heuristic_profile(text)
    assert p.name == "Arjun Mehta"
    assert p.email == "arjun.mehta@example.com"
    assert p.years_experience == 9
    assert [(j.start, j.end) for j in p.employment_history] == [("2019", "present"), ("2015", "2019")]
    assert "Razorpay" in p.employment_history[0].title
    assert {"Python", "REST API design", "Terraform"} <= set(p.skills)
    assert p.links == ["linkedin.com/in/arjunmehta"]


def test_docx_text_includes_tables_and_headers():
    from docx import Document

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Meera Nair · meera@example.com"
    doc.add_paragraph("Senior QA Engineer")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Selenium, Python, JIRA"
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text("resume.docx", buf.getvalue())
    assert "Senior QA Engineer" in text
    assert "Skills | Selenium, Python, JIRA" in text
    assert "meera@example.com" in text
    # Paragraphs keep their order relative to tables.
    assert text.index("Senior QA Engineer") < text.index("Selenium")


def test_existing_docx_sample_still_reads():
    text = extract_text("r.docx", (TEST_DATA / "05-docx-upload.docx").read_bytes())
    assert len(text) > 100


def test_structured_profile_is_saved_and_shown(client):
    text = (TEST_DATA / "01-strong-backend-fit.txt").read_text(encoding="utf-8")
    cand = client.post("/api/candidates", json={"resume_text": text}).json()
    detail = client.get(f"/api/candidates/{cand['id']}").json()
    history = detail["profile"]["employment_history"]
    assert len(history) == 2 and history[0]["end"] == "present"
    assert detail["profile"]["years_from_dates"] is not None


def test_text_that_is_not_a_resume_is_refused(client):
    letter = "Dear hiring team,\nI am writing to say how much I admire your company and would love to hear from you.\nThanks!"
    r = client.post("/api/candidates", json={"resume_text": letter})
    assert r.status_code == 422
    assert "Is this a resume?" in r.json()["detail"]
    assert client.get("/api/candidates").json() == []
